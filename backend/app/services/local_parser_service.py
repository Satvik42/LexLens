"""Local parsing fallback (no bounding boxes) producing the same normalized structure as Document AI.

Used when Document AI is not configured (development, tests, offline demos).
"""

import logging
from io import BytesIO

from app.schemas.document import NormalizedBlock, NormalizedPage, ParsedDocument
from app.services.text_utils import make_block_id, paragraphs_from_text

logger = logging.getLogger(__name__)

PAGE_BREAK = "\f"
_PARAGRAPHS_PER_SYNTHETIC_PAGE = 12


class UnsupportedDocumentError(Exception):
    """Raised when a document has no extractable text."""


def _page_from_paragraphs(page_number: int, paragraphs: list[str]) -> NormalizedPage:
    blocks: list[NormalizedBlock] = []
    cursor = 0
    for index, paragraph in enumerate(paragraphs):
        start = cursor
        end = start + len(paragraph)
        blocks.append(NormalizedBlock(block_id=make_block_id(page_number, index), text=paragraph, start=start, end=end))
        cursor = end + 2
    return NormalizedPage(page_number=page_number, text="\n\n".join(paragraphs), blocks=blocks)


def _pages_from_texts(page_texts: list[str]) -> list[NormalizedPage]:
    pages = []
    for number, text in enumerate(page_texts, start=1):
        paragraphs = paragraphs_from_text(text)
        if paragraphs:
            pages.append(_page_from_paragraphs(len(pages) + 1, paragraphs))
    return pages


def _parse_pdf(content: bytes) -> list[NormalizedPage]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    if reader.is_encrypted:
        raise UnsupportedDocumentError("Encrypted PDFs are not supported")
    return _pages_from_texts([(page.extract_text() or "") for page in reader.pages])


def _parse_docx(content: bytes) -> list[NormalizedPage]:
    from docx import Document as DocxDocument

    document = DocxDocument(BytesIO(content))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return _paginate(paragraphs)


def _parse_txt(content: bytes) -> list[NormalizedPage]:
    text = content.decode("utf-8", errors="replace")
    if PAGE_BREAK in text:
        return _pages_from_texts(text.split(PAGE_BREAK))
    return _paginate(paragraphs_from_text(text))


def _paginate(paragraphs: list[str]) -> list[NormalizedPage]:
    """Formats without pages (DOCX/TXT) are grouped into synthetic pages so evidence still has a page reference."""
    pages = []
    for offset in range(0, len(paragraphs), _PARAGRAPHS_PER_SYNTHETIC_PAGE):
        chunk = paragraphs[offset : offset + _PARAGRAPHS_PER_SYNTHETIC_PAGE]
        pages.append(_page_from_paragraphs(len(pages) + 1, chunk))
    return pages


_PARSERS = {
    "application/pdf": _parse_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _parse_docx,
    "text/plain": _parse_txt,
}


class LocalParserService:
    def parse(self, content: bytes, mime_type: str) -> ParsedDocument:
        parser = _PARSERS.get(mime_type)
        if parser is None:
            raise UnsupportedDocumentError(f"No parser for {mime_type}")
        try:
            pages = parser(content)
        except UnsupportedDocumentError:
            raise
        except Exception as exc:  # corrupt files are expected user input, not a server fault
            raise UnsupportedDocumentError("The document could not be parsed") from exc
        if not pages:
            raise UnsupportedDocumentError("No extractable text found")
        return ParsedDocument(pages=pages, parser="local")
